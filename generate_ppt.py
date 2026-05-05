from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN

def create_ppt():
    prs = Presentation()
    
    # 定义一些辅助函数
    def add_title_slide(title, subtitle):
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        title_box = slide.shapes.title
        subtitle_box = slide.placeholders[1]
        
        title_box.text = title
        subtitle_box.text = subtitle
        
    def add_bullet_slide(title, content_lines):
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        title_box = slide.shapes.title
        body_box = slide.placeholders[1]
        
        title_box.text = title
        tf = body_box.text_frame
        
        for i, line in enumerate(content_lines):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = line
            p.level = 0
            if line.startswith("  -"):
                p.level = 1
                p.text = line[3:].strip()
            elif line.startswith("- "):
                p.text = line[2:].strip()

    # Slide 1
    add_title_slide("探索之书 (The Book of Exploration)", "在当下，看见最深处的自己\n\n意料之外，情理之中 —— 你的两分钟身心重置指南")
    
    # Slide 2
    add_bullet_slide("痛点洞察：身心失联", [
        "- 焦虑、失眠、拖延的背后：人与自身当下的“断连”。",
        "- 传统心理干预的死局：大道理泛滥，处于情绪内耗中的人最反感说教。",
        "- 核心痛点：我们不仅需要情绪的创可贴，更需要一面能映照潜意识的镜子。"
    ])
    
    # Slide 3
    add_bullet_slide("我们的解法：东方修行哲学", [
        "- 不给说教，只给微行动：2 分钟内可完成的意料之外的动作。",
        "- 暗合“天、地、人”修行智慧：",
        "  - ☁️ 天：调息与念头（如颅腔共鸣，向下剥离思绪）",
        "  - 🌍 地：扎根与身体连接（如脚趾抓地，向上引导重力）",
        "  - 🧍 人：情绪与向内觉察（如轻捏眉心，释放社交紧绷）"
    ])
    
    # Slide 4
    add_bullet_slide("Aha Moment：意料之外，情理之中", [
        "- 案例：睡不着，脑子停不下来",
        "- 行动：用脚趾抓床单（属于“地”的连接）。",
        "- 科学原理：注意力强行从大脑拉回地面，激活副交感神经。",
        "- 结果：大脑停机，真正回到“当下”。"
    ])
    
    # Slide 5
    add_bullet_slide("深层洞察：看见微习惯", [
        "- 每一次选择，都是潜意识的表达：",
        "  - 偏好“地”：可能长期缺乏安全感，需要物理连接。",
        "  - 偏好“天”：可能思维过度活跃，需要清空缓存。",
        "- 数据挖掘：通过自然语言提问与选择偏好，穿透表象的失眠，看见深层的职场恐惧或关系焦虑。"
    ])
    
    # Slide 6
    add_bullet_slide("商业闭环：从工具到转化", [
        "- 1. 引流层 (免费)：2 分钟高频微行动，建立信任。",
        "- 2. 留存与洞察：沉淀行为数据，定期生成「内在探索报告」。",
        "- 3. 精准商业转化：顺理成章地推荐深度服务。",
        "  - 例如：定制化冥想课程、线下禅修营、一对一专业心理咨询。"
    ])
    
    # Slide 7
    add_title_slide("愿景", "让每一个问题，\n都成为一次向内探索的修行。")
    
    # Slide 8
    add_title_slide("谢谢大家！", "翻开探索之书，遇见当下的自己。")

    prs.save('docs/探索之书_路演.pptx')
    print("PPTX generated successfully!")

if __name__ == '__main__':
    create_ppt()
